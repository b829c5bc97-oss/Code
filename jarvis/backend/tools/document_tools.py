"""
Document generation tools — Phase 7.

Covers the formats listed in the project spec: TXT/Markdown/CSV need only
the standard library; DOCX/PPTX/XLSX/PDF use the well-established
`python-docx` / `python-pptx` / `openpyxl` / `fpdf2` libraries rather than
reinventing file formats. A relative `path` resolves under
`Settings.workspace_dir`; an absolute path is used as-is (so "make a report
on my Desktop" still works).
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

from backend.core.config import get_settings


class DocumentError(RuntimeError):
    """Raised when a document can't be created."""


def _resolve(path: str) -> Path:
    p = Path(path).expanduser()
    if not p.is_absolute():
        p = get_settings().workspace_path / p
    p.parent.mkdir(parents=True, exist_ok=True)
    return p.resolve()


def create_txt(path: str, content: str) -> str:
    p = _resolve(path)
    p.write_text(content, encoding="utf-8")
    return str(p)


def create_markdown(path: str, content: str) -> str:
    p = _resolve(path)
    p.write_text(content, encoding="utf-8")
    return str(p)


def create_csv(path: str, rows: list[list[str]]) -> str:
    p = _resolve(path)
    with p.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)
    return str(p)


def create_docx(path: str, title: str, paragraphs: list[str]) -> str:
    try:
        import docx
    except ImportError as exc:
        raise DocumentError("The 'python-docx' package is not installed.") from exc

    document = docx.Document()
    document.add_heading(title, level=1)
    for para in paragraphs:
        document.add_paragraph(para)
    p = _resolve(path)
    document.save(p)
    return str(p)


def create_pptx(path: str, title: str, slides: list[dict]) -> str:
    """`slides`: list of {"title": str, "bullets": list[str]}."""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise DocumentError("The 'python-pptx' package is not installed.") from exc

    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    bullet_layout = prs.slide_layouts[1]

    title_slide = prs.slides.add_slide(title_layout)
    title_slide.shapes.title.text = title

    for slide_data in slides:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = slide_data.get("title", "")
        body = slide.placeholders[1].text_frame
        bullets = slide_data.get("bullets", [])
        if bullets:
            body.text = bullets[0]
            for bullet in bullets[1:]:
                body.add_paragraph().text = bullet

    p = _resolve(path)
    prs.save(p)
    return str(p)


def create_xlsx(path: str, rows: list[list], sheet_name: str = "Sheet1") -> str:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise DocumentError("The 'openpyxl' package is not installed.") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    p = _resolve(path)
    wb.save(p)
    return str(p)


def create_pdf(path: str, title: str, paragraphs: list[str]) -> str:
    try:
        from fpdf import FPDF
    except ImportError as exc:
        raise DocumentError("The 'fpdf2' package is not installed.") from exc

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, title)
    pdf.ln(4)
    pdf.set_font("Helvetica", size=12)
    for para in paragraphs:
        pdf.multi_cell(0, 8, para)
        pdf.ln(2)
    p = _resolve(path)
    pdf.output(str(p))
    return str(p)


def read_csv(path: str) -> list[list[str]]:
    p = Path(path).expanduser().resolve()
    if not p.exists():
        raise DocumentError(f"{p} does not exist.")
    with p.open(newline="", encoding="utf-8") as f:
        return list(csv.reader(f))


def register_document_tools(registry) -> None:  # noqa: ANN001 - avoid import cycle for typing only
    from backend.core.permissions import PermissionLevel
    from backend.tools.registry import Tool, ToolResult

    async def txt(path: str, content: str, **_: object) -> ToolResult:
        try:
            saved = create_txt(path, content)
        except (DocumentError, OSError) as exc:
            return ToolResult(success=False, error=str(exc))
        return ToolResult(success=True, output=f"Saved {saved}.")

    async def markdown(path: str, content: str, **_: object) -> ToolResult:
        try:
            saved = create_markdown(path, content)
        except (DocumentError, OSError) as exc:
            return ToolResult(success=False, error=str(exc))
        return ToolResult(success=True, output=f"Saved {saved}.")

    async def csv_doc(path: str, rows: list[list[str]], **_: object) -> ToolResult:
        try:
            saved = create_csv(path, rows)
        except (DocumentError, OSError) as exc:
            return ToolResult(success=False, error=str(exc))
        return ToolResult(success=True, output=f"Saved {saved}.")

    async def docx_doc(path: str, title: str, paragraphs: list[str], **_: object) -> ToolResult:
        try:
            saved = create_docx(path, title, paragraphs)
        except (DocumentError, OSError) as exc:
            return ToolResult(success=False, error=str(exc))
        return ToolResult(success=True, output=f"Saved {saved}.")

    async def pptx_doc(path: str, title: str, slides: list[dict], **_: object) -> ToolResult:
        try:
            saved = create_pptx(path, title, slides)
        except (DocumentError, OSError) as exc:
            return ToolResult(success=False, error=str(exc))
        return ToolResult(success=True, output=f"Saved {saved}.")

    async def xlsx_doc(path: str, rows: list[list], sheet_name: str = "Sheet1", **_: object) -> ToolResult:
        try:
            saved = create_xlsx(path, rows, sheet_name)
        except (DocumentError, OSError) as exc:
            return ToolResult(success=False, error=str(exc))
        return ToolResult(success=True, output=f"Saved {saved}.")

    async def pdf_doc(path: str, title: str, paragraphs: list[str], **_: object) -> ToolResult:
        try:
            saved = create_pdf(path, title, paragraphs)
        except (DocumentError, OSError) as exc:
            return ToolResult(success=False, error=str(exc))
        return ToolResult(success=True, output=f"Saved {saved}.")

    registry.register(
        Tool(
            name="documents.create_txt",
            description="Create a plain text file.",
            parameters={
                "type": "object",
                "properties": {"path": {"type": "string"}, "content": {"type": "string"}},
                "required": ["path", "content"],
            },
            permission=PermissionLevel.MEDIUM,
            execute=txt,
        )
    )
    registry.register(
        Tool(
            name="documents.create_markdown",
            description="Create a Markdown (.md) file.",
            parameters={
                "type": "object",
                "properties": {"path": {"type": "string"}, "content": {"type": "string"}},
                "required": ["path", "content"],
            },
            permission=PermissionLevel.MEDIUM,
            execute=markdown,
        )
    )
    registry.register(
        Tool(
            name="documents.create_csv",
            description="Create a CSV file from rows of values.",
            parameters={
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "rows": {"type": "array", "items": {"type": "array"}},
                },
                "required": ["path", "rows"],
            },
            permission=PermissionLevel.MEDIUM,
            execute=csv_doc,
        )
    )
    registry.register(
        Tool(
            name="documents.create_docx",
            description="Create a Word (.docx) document with a title and paragraphs.",
            parameters={
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "title": {"type": "string"},
                    "paragraphs": {"type": "array", "items": {"type": "string"}},
                },
                "required": ["path", "title", "paragraphs"],
            },
            permission=PermissionLevel.MEDIUM,
            execute=docx_doc,
        )
    )
    registry.register(
        Tool(
            name="documents.create_pptx",
            description=(
                "Create a PowerPoint (.pptx) presentation. `slides` is a list of "
                '{"title": str, "bullets": [str, ...]}.'
            ),
            parameters={
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "title": {"type": "string"},
                    "slides": {"type": "array", "items": {"type": "object"}},
                },
                "required": ["path", "title", "slides"],
            },
            permission=PermissionLevel.MEDIUM,
            execute=pptx_doc,
        )
    )
    registry.register(
        Tool(
            name="documents.create_xlsx",
            description="Create an Excel (.xlsx) spreadsheet from rows of values.",
            parameters={
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "rows": {"type": "array", "items": {"type": "array"}},
                    "sheet_name": {"type": "string", "default": "Sheet1"},
                },
                "required": ["path", "rows"],
            },
            permission=PermissionLevel.MEDIUM,
            execute=xlsx_doc,
        )
    )
    registry.register(
        Tool(
            name="documents.create_pdf",
            description="Create a simple PDF document with a title and paragraphs.",
            parameters={
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "title": {"type": "string"},
                    "paragraphs": {"type": "array", "items": {"type": "string"}},
                },
                "required": ["path", "title", "paragraphs"],
            },
            permission=PermissionLevel.MEDIUM,
            execute=pdf_doc,
        )
    )
